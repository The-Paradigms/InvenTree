import openpyxl

wb = openpyxl.load_workbook(r"C:\Users\srini\Downloads\items_20260508.xlsx", read_only=True, data_only=True)
ws = wb.active
headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]

print("=== CATEGORY ROWS ===")
count = 0
for row in ws.iter_rows(min_row=2, values_only=True):
    r = dict(zip(headers, row))
    if r["TypeOfItem"] == "Category":
        print("  Name=" + repr(r["Name"]) + "  Parent=" + repr(r["Category"]))
        count += 1
        if count >= 30:
            break

print()
print("=== SAMPLE INVENTORY ROWS ===")
count = 0
for row in ws.iter_rows(min_row=2, values_only=True):
    r = dict(zip(headers, row))
    if r["TypeOfItem"] == "Inventory Item":
        print("  Name=" + repr(r["Name"]))
        print("    SKU=" + repr(r["SKU"]) + "  Barcode=" + repr(r["Barcode"]) + "  Category=" + repr(r["Category"]))
        print("    Vendor=" + repr(r["PreferredVendor"]) + "  VPN=" + repr(r["VendorPartNumber"]))
        print("    QOH=" + str(r["QuantityOnHand"]) + "  Cost=" + str(r["Cost"]) + "  SalesPrice=" + str(r["SalesPrice"]))
        print("    Desc=" + repr(str(r["SalesDescription"])[:80]))
        print()
        count += 1
        if count >= 6:
            break

wb.close()
