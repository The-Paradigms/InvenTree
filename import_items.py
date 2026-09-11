"""
Import items from items_20260508.xlsx into InvenTree.

Run from repo root:
    python import_items.py

Creates:
  - PartCategory   (54 flat categories)
  - Part           (8944 inventory items)
  - Company        (23 suppliers, is_supplier=True)
  - SupplierPart   (one per part that has a preferred vendor)
  - StockItem      (one per part with QuantityOnHand > 0)
  - StockLocation  (single default "Warehouse" location)
"""

import os
import sys
import django
from decimal import Decimal, InvalidOperation

EXCEL_PATH = r"C:\Users\srini\Downloads\items_20260508.xlsx"

BACKEND_DIR = os.path.join(os.path.dirname(__file__), "src", "backend", "InvenTree")
os.chdir(BACKEND_DIR)
sys.path.insert(0, BACKEND_DIR)
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "InvenTree.settings")
django.setup()

# ----- imports after django.setup() -----
import openpyxl
from part.models import Part, PartCategory
from company.models import Company, SupplierPart
from stock.models import StockItem, StockLocation

# ------------------------------------------------------------------ helpers

def to_decimal(val, default=Decimal("0")):
    try:
        return Decimal(str(val)) if val else default
    except InvalidOperation:
        return default

def clean(val, default=""):
    if val is None:
        return default
    s = str(val).strip()
    # Strip replacement character from broken encoding
    s = s.replace("�", "").strip()
    return s if s else default

# ------------------------------------------------------------------ load workbook

print("Loading workbook...")
wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
ws = wb.active
headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
rows = [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]
wb.close()
print(f"  {len(rows)} rows loaded")

# ------------------------------------------------------------------ 1. categories

print("\n[1/5] Creating PartCategories...")

# Collect all category names: from Category rows + names referenced by Inventory Items
category_names = set()
for r in rows:
    if r["TypeOfItem"] == "Category":
        name = clean(r["Name"])
        if name:
            category_names.add(name)
for r in rows:
    if r["TypeOfItem"] == "Inventory Item":
        name = clean(r["Category"])
        if name:
            category_names.add(name)

cat_map = {}  # name -> PartCategory instance
created_cats = 0
for name in sorted(category_names):
    obj, created = PartCategory.objects.get_or_create(
        name=name,
        defaults={"description": name},
    )
    cat_map[name] = obj
    if created:
        created_cats += 1

print(f"  {created_cats} created, {len(cat_map) - created_cats} already existed  ({len(cat_map)} total)")

# ------------------------------------------------------------------ 2. suppliers (Company)

print("\n[2/5] Creating Supplier companies...")

vendor_names = set()
for r in rows:
    v = clean(r["PreferredVendor"])
    if v:
        vendor_names.add(v)

vendor_map = {}  # name -> Company instance
created_vendors = 0
for name in sorted(vendor_names):
    obj, created = Company.objects.get_or_create(
        name=name,
        defaults={
            "is_supplier": True,
            "is_manufacturer": False,
            "is_customer": False,
            "description": name,
        },
    )
    # Make sure existing company is marked as supplier
    if not obj.is_supplier:
        obj.is_supplier = True
        obj.save(update_fields=["is_supplier"])
    vendor_map[name] = obj
    if created:
        created_vendors += 1

print(f"  {created_vendors} created, {len(vendor_map) - created_vendors} already existed  ({len(vendor_map)} total)")

# ------------------------------------------------------------------ 3. parts

print("\n[3/5] Creating Parts...")

inventory_rows = [r for r in rows if r["TypeOfItem"] == "Inventory Item"]
part_map = {}  # row Name -> Part instance
created_parts = 0
skipped_parts = 0
errors = []

for i, r in enumerate(inventory_rows):
    name = clean(r["Name"])
    if not name:
        skipped_parts += 1
        continue

    description = clean(r["SalesDescription"]) or clean(r["PurchaseDescription"]) or name
    # Django field max_length=250
    description = description[:250]

    category = cat_map.get(clean(r["Category"]))
    ipn = clean(r["SKU"])
    barcode = clean(r["Barcode"])
    notes = clean(r["Notes"])
    minimum_stock = to_decimal(r["ReorderLevel"])
    maximum_stock = to_decimal(r["MaxStockLevel"]) if r["MaxStockLevel"] else None
    active = not bool(r["Archived"])

    defaults = {
        "description": description,
        "category": category,
        "notes": notes,
        "minimum_stock": minimum_stock,
        "active": active,
        "purchaseable": bool(r["PurchasingItem"]),
        "salable": bool(r["SalesItem"]),
        "component": bool(r["ManufacturingItem"]),
    }
    if ipn:
        defaults["IPN"] = ipn[:100]
    if barcode:
        defaults["barcode_data"] = barcode
    if maximum_stock:
        defaults["maximum_stock"] = maximum_stock

    try:
        obj, created = Part.objects.get_or_create(name=name, defaults=defaults)
        part_map[name] = obj
        if created:
            created_parts += 1
    except Exception as e:
        errors.append(f"Part '{name}': {e}")
        skipped_parts += 1

    if (i + 1) % 500 == 0:
        print(f"  ... {i + 1}/{len(inventory_rows)}")

print(f"  {created_parts} created, {len(part_map) - created_parts} already existed, {skipped_parts} skipped")
if errors:
    print(f"  Errors ({len(errors)}):")
    for e in errors[:10]:
        print(f"    {e}")

# ------------------------------------------------------------------ 4. supplier parts

print("\n[4/5] Creating SupplierParts...")
created_sp = 0
skipped_sp = 0
sp_errors = []

for r in inventory_rows:
    name = clean(r["Name"])
    vendor_name = clean(r["PreferredVendor"])
    if not name or not vendor_name:
        continue

    part = part_map.get(name)
    supplier = vendor_map.get(vendor_name)
    if not part or not supplier:
        continue

    # SKU is required; fall back to part name if vendor part number is empty
    vpn = clean(r["VendorPartNumber"]) or name
    purchase_desc = clean(r["PurchaseDescription"])

    defaults = {
        "description": purchase_desc[:250] if purchase_desc else "",
    }

    try:
        obj, created = SupplierPart.objects.get_or_create(
            part=part,
            supplier=supplier,
            SKU=vpn,
            defaults=defaults,
        )
        if created:
            created_sp += 1
    except Exception as e:
        sp_errors.append(f"SupplierPart '{name}' / '{vendor_name}': {e}")
        skipped_sp += 1

print(f"  {created_sp} created, skipped {skipped_sp}")
if sp_errors:
    print(f"  Errors ({len(sp_errors)}):")
    for e in sp_errors[:10]:
        print(f"    {e}")

# ------------------------------------------------------------------ 5. stock items

print("\n[5/5] Creating StockItems...")

default_location, _ = StockLocation.objects.get_or_create(
    name="Warehouse",
    defaults={"description": "Default warehouse location"},
)

created_stock = 0
skipped_stock = 0
stock_errors = []

for r in inventory_rows:
    qty = to_decimal(r["QuantityOnHand"])
    if qty <= 0:
        continue

    name = clean(r["Name"])
    part = part_map.get(name)
    if not part:
        skipped_stock += 1
        continue

    cost = to_decimal(r["Cost"])

    # Only create if no stock item already exists for this part
    if StockItem.objects.filter(part=part).exists():
        continue

    defaults = {
        "quantity": qty,
        "location": default_location,
    }
    if cost > 0:
        defaults["purchase_price"] = cost
        defaults["purchase_price_currency"] = "USD"

    try:
        StockItem.objects.create(part=part, **defaults)
        created_stock += 1
    except Exception as e:
        stock_errors.append(f"StockItem '{name}': {e}")
        skipped_stock += 1

print(f"  {created_stock} created, {skipped_stock} skipped")
if stock_errors:
    print(f"  Errors ({len(stock_errors)}):")
    for e in stock_errors[:10]:
        print(f"    {e}")

# ------------------------------------------------------------------ summary

print("\n=== DONE ===")
print(f"  Categories:    {PartCategory.objects.count()} total")
print(f"  Parts:         {Part.objects.count()} total")
print(f"  Companies:     {Company.objects.filter(is_supplier=True).count()} suppliers")
print(f"  SupplierParts: {SupplierPart.objects.count()} total")
print(f"  StockItems:    {StockItem.objects.count()} total")
